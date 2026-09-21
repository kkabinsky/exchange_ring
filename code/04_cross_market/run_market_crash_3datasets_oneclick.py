from __future__ import annotations

import argparse
import json
import math
import os
import shutil
import subprocess
import sys
import time
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# --- replication package (added): data in <repo>/input, results in <repo>/output/04_cross_market ---
import os as _rp_os, sys as _rp_sys
from pathlib import Path as _RpPath
REPO_ROOT = _RpPath(__file__).resolve().parents[2]
REPO_INPUT = REPO_ROOT / "input"
REPO_OUTPUT = REPO_ROOT / "output" / "04_cross_market"
REPO_OUTPUT.mkdir(parents=True, exist_ok=True)
_rp_sys.path.insert(0, str(REPO_ROOT / "code" / "common"))
# ------------------------------------------------------------------------------------------
BASE = REPO_OUTPUT  # results folder (was the script folder)
SINGLE_JOB = Path(__file__).resolve().parent / "run_single_fanogan_job.py"
MASTER_REPORT = BASE / "market_crash_3datasets_3crashes_summary.xlsx"
WINDOW_SIZE = 32
METHODS = ["cosine", "exponential", "arctan", "arccosh"]

DATASETS = [
    {
        "dataset": "S&P500",
        "code": "SP500_U500",
        "input": REPO_INPUT / "u500_1day.xlsx",
        "price_col": "closed",
        "date_col": "date",
        "sheet": "0",
    },
    {
        "dataset": "DAX INDEX",
        "code": "DAX_INDEX",
        "input": REPO_INPUT / "dax_1day.xlsx",
        "price_col": "closed",
        "date_col": "date",
        "sheet": "0",
    },
    {
        "dataset": "DJIA",
        "code": "DJIA_U30",
        "input": REPO_INPUT / "u30_1day.xlsx",
        "price_col": "closed",
        "date_col": "date",
        "sheet": "0",
    },
]

SCENARIOS = [
    {
        "scenario": "COVID-19 market crash",
        "code": "COVID19_market_crash",
        "train_start": "2013-11-01",
        "train_end": "2019-12-05",
        "test_start": "2019-12-06",
        "test_end": "2020-03-31",
        # Documented crash sub-window inside the test span (S&P500 peak 2020-02-19,
        # trough 2020-03-23). Windows before the onset are the test-normal class.
        "crash_onset": "2020-02-20",
        "crash_end_actual": "2020-03-23",
    },
    {
        "scenario": "Russia-Ukraine war crash",
        "code": "Russia_Ukraine_war_crash",
        "train_start": "2013-11-01",
        "train_end": "2018-02-10",
        "test_start": "2021-10-28",
        "test_end": "2022-02-24",
        # Invasion 2022-02-24; war-risk selloff intensified from mid-Feb 2022.
        "crash_onset": "2022-02-11",
        "crash_end_actual": "2022-02-24",
    },
    {
        "scenario": "Chinese real asset market crash",
        "code": "Chinese_real_asset_market_crash",
        "train_start": "2013-11-01",
        "train_end": "2018-02-10",
        "test_start": "2023-05-01",
        "test_end": "2023-10-31",
        # Country Garden/Evergrande property crisis deepened Aug-Oct 2023.
        # test_start is set earlier (2023-05-01) to include a pre-crash normal
        # baseline, consistent with the other scenarios.
        "crash_onset": "2023-08-07",
        "crash_end_actual": "2023-10-31",
        "note": "The supplied table ended at 2023/10/, so this runner uses 2023-10-31.",
    },
]


def parse_dates(series: pd.Series) -> pd.Series:
    text = series.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    parsed = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")

    be_mask = text.str.fullmatch(r"(24|25|26)\d{6}")
    if be_mask.any():
        be_text = text[be_mask]
        be_year = pd.to_numeric(be_text.str.slice(0, 4), errors="coerce") - 543
        gregorian = (
            be_year.astype("Int64").astype(str)
            + "-"
            + be_text.str.slice(4, 6)
            + "-"
            + be_text.str.slice(6, 8)
        )
        parsed.loc[be_mask] = pd.to_datetime(gregorian, format="%Y-%m-%d", errors="coerce")

    for formatted_text, fmt in [
        (text.str.zfill(6), "%d%m%y"),
        (text.str.zfill(8), "%d%m%Y"),
        (text.str.zfill(8), "%Y%m%d"),
    ]:
        parsed = parsed.fillna(pd.to_datetime(formatted_text, format=fmt, errors="coerce"))

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        parsed = parsed.fillna(pd.to_datetime(text, errors="coerce", dayfirst=True))
    return parsed


def find_external_fanogan_runs() -> list[dict]:
    ps = f"""
    Get-CimInstance Win32_Process |
      Where-Object {{
        $_.Name -in @('python.exe','py.exe') -and
        $_.ProcessId -ne {os.getpid()} -and
        $_.CommandLine -match 'fanogan_four_gaf_compare_2|run_us1m|run_u500|run_single_fanogan_job|run_market_crash_3datasets'
      }} |
      Select-Object ProcessId,CreationDate,CommandLine |
      ConvertTo-Json -Compress
    """
    try:
        out = subprocess.check_output(
            ["powershell", "-NoProfile", "-Command", ps],
            text=True,
            stderr=subprocess.DEVNULL,
        ).strip()
    except Exception:
        return []
    if not out:
        return []
    parsed = json.loads(out)
    return parsed if isinstance(parsed, list) else [parsed]


def wait_for_external_fanogan_runs() -> None:
    while True:
        runs = find_external_fanogan_runs()
        if not runs:
            return
        pids = ", ".join(str(run.get("ProcessId")) for run in runs)
        print(f"[WAIT] Another f-AnoGAN job is running: PID {pids}. Waiting before starting this batch.", flush=True)
        time.sleep(60)


def load_dataset_frame(dataset: dict) -> pd.DataFrame:
    df = pd.read_excel(dataset["input"], sheet_name=int(dataset["sheet"]))
    cols_lower = {str(c).lower().strip(): c for c in df.columns}
    price_col = cols_lower[dataset["price_col"].lower()]
    date_col = cols_lower[dataset["date_col"].lower()]

    out = pd.DataFrame(
        {
            "Date": parse_dates(df[date_col]),
            "price": pd.to_numeric(df[price_col], errors="coerce"),
        }
    )
    return out.dropna().sort_values("Date").drop_duplicates("Date", keep="last").reset_index(drop=True)


def compute_job_config(dataset: dict, scenario: dict) -> dict:
    frame = load_dataset_frame(dataset)
    starts = np.arange(len(frame) - WINDOW_SIZE + 1)
    if len(starts) <= 0:
        raise RuntimeError(f"Not enough rows for {dataset['code']}")

    start_dates = frame.loc[starts, "Date"].reset_index(drop=True)
    train_end = pd.Timestamp(scenario["train_end"])
    test_start = pd.Timestamp(scenario["test_start"])
    test_end = pd.Timestamp(scenario["test_end"])

    test_mask = (start_dates >= test_start) & (start_dates <= test_end)
    test_positions = np.where(test_mask.to_numpy())[0]
    if len(test_positions) == 0:
        raise RuntimeError(f"No test windows for {dataset['code']} {scenario['code']}")

    # Crash region = the documented historical crash sub-window, NOT the whole
    # test span. This keeps test-normal windows (before the crash onset) in the
    # test set so AUC/F1/DM are well defined (test contains both classes).
    crash_onset_ts = pd.Timestamp(scenario.get("crash_onset", scenario["test_start"]))
    crash_actual_end_ts = pd.Timestamp(scenario.get("crash_end_actual", scenario["test_end"]))
    crash_start_candidates = frame.index[frame["Date"] >= crash_onset_ts].tolist()
    crash_end_candidates = frame.index[frame["Date"] > crash_actual_end_ts].tolist()
    if not crash_start_candidates:
        raise RuntimeError(f"No crash start date for {dataset['code']} {scenario['code']}")
    crash_start = int(crash_start_candidates[0])
    crash_end = int(crash_end_candidates[0]) if crash_end_candidates else int(len(frame))
    if crash_end <= crash_start:
        raise RuntimeError(f"Invalid crash index for {dataset['code']} {scenario['code']}")

    labels = np.array([1 if ((int(s) < crash_end) and (int(s) + WINDOW_SIZE > crash_start)) else 0 for s in starts])
    train_mask = start_dates <= train_end

    actual_train_start = frame.loc[frame["Date"] >= pd.Timestamp(scenario["train_start"]), "Date"]
    actual_train_start = actual_train_start.iloc[0] if not actual_train_start.empty else frame["Date"].iloc[0]
    train_last = frame.loc[frame["Date"] <= train_end, "Date"]
    actual_train_end = train_last.iloc[-1] if not train_last.empty else pd.NaT

    return {
        "dataset": dataset["dataset"],
        "dataset_code": dataset["code"],
        "scenario": scenario["scenario"],
        "scenario_code": scenario["code"],
        "stock_name": f"{dataset['code']}_{scenario['code']}",
        "input": str(dataset["input"]),
        "train_start_table": scenario["train_start"],
        "train_end_table": scenario["train_end"],
        "test_start_table": scenario["test_start"],
        "test_end_table": scenario["test_end"],
        "actual_data_start": frame["Date"].iloc[0].date().isoformat(),
        "actual_data_end": frame["Date"].iloc[-1].date().isoformat(),
        "actual_train_start": actual_train_start.date().isoformat(),
        "actual_train_end": "" if pd.isna(actual_train_end) else actual_train_end.date().isoformat(),
        "actual_test_start": frame.loc[int(starts[test_positions[0]]), "Date"].date().isoformat(),
        "actual_test_end": frame.loc[int(starts[test_positions[-1]]), "Date"].date().isoformat(),
        "forecast_days": int(len(test_positions)),
        "crash_start_index": crash_start,
        "crash_end_index_exclusive": crash_end,
        "crash_onset_date": scenario.get("crash_onset", scenario["test_start"]),
        "crash_end_date": scenario.get("crash_end_actual", scenario["test_end"]),
        "n_price_rows": int(len(frame)),
        "n_windows_total": int(len(starts)),
        "n_train_windows": int(train_mask.sum()),
        "n_train_normal_windows": int((train_mask.to_numpy() & (labels == 0)).sum()),
        "n_test_windows": int(len(test_positions)),
        "n_test_crash_windows": int(labels[test_positions].sum()),
        "n_test_normal_windows": int((labels[test_positions] == 0).sum()),
        "note": scenario.get("note", ""),
    }


def normal_cdf(x: float) -> float:
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def dm_tests(scores: pd.DataFrame) -> pd.DataFrame:
    key_cols = ["window_start", "window_end"]
    wide = None
    for method in METHODS:
        df = scores.loc[scores["method"] == method, key_cols + ["label_0normal_1crash", "predicted_label"]].copy()
        df[f"loss_{method}"] = (df["label_0normal_1crash"] - df["predicted_label"]) ** 2
        df = df[key_cols + [f"loss_{method}"]]
        wide = df if wide is None else wide.merge(df, on=key_cols, how="inner")

    rows = []
    for other in ["cosine", "arctan", "arccosh"]:
        d = wide[f"loss_{other}"] - wide["loss_exponential"]
        t = int(d.shape[0])
        mean_diff = float(d.mean())
        var = float(d.var(ddof=1)) if t > 1 else float("nan")
        if t > 1 and var > 0:
            dm_stat = mean_diff / math.sqrt(var / t)
            p_value = 2.0 * (1.0 - normal_cdf(abs(dm_stat)))
        elif mean_diff == 0:
            dm_stat = 0.0
            p_value = 1.0
        else:
            dm_stat = math.copysign(float("inf"), mean_diff)
            p_value = 0.0

        exp_loss = float(wide["loss_exponential"].mean())
        other_loss = float(wide[f"loss_{other}"].mean())
        rows.append(
            {
                "comparison": f"exponential_vs_{other}",
                "T": t,
                "loss_exponential": exp_loss,
                "loss_other": other_loss,
                "mean_loss_diff_other_minus_exp": mean_diff,
                "dm_statistic": dm_stat,
                "p_value": p_value,
                "winner_lower_loss": "exponential" if exp_loss < other_loss else other if other_loss < exp_loss else "tie",
                "interpretation": "positive DM favors exponential" if mean_diff > 0 else "negative DM favors other" if mean_diff < 0 else "same loss",
            }
        )
    return pd.DataFrame(rows)


def exponential_outperformance_summary(perf_all: pd.DataFrame, dm_all: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    detail_rows = []
    group_cols = ["dataset_code", "dataset", "scenario_code", "scenario"]
    for keys, group in perf_all.groupby(group_cols, dropna=False):
        dataset_code, dataset, scenario_code, scenario = keys
        ranked = group.sort_values(["f1", "auc", "accuracy"], ascending=False).reset_index(drop=True)
        exp_rows = ranked.loc[ranked["method"] == "exponential"]
        if exp_rows.empty:
            continue
        exp = exp_rows.iloc[0]
        best = ranked.iloc[0]
        exp_rank = int(ranked.index[ranked["method"] == "exponential"][0]) + 1

        dm_group = dm_all[
            (dm_all["dataset_code"] == dataset_code)
            & (dm_all["scenario_code"] == scenario_code)
        ].copy()
        significant_wins = int(((dm_group["winner_lower_loss"] == "exponential") & (dm_group["p_value"] < 0.05)).sum())
        dm_wins = int((dm_group["winner_lower_loss"] == "exponential").sum())
        significant_losses = int(((dm_group["winner_lower_loss"] != "exponential") & (dm_group["winner_lower_loss"] != "tie") & (dm_group["p_value"] < 0.05)).sum())

        detail_rows.append(
            {
                "dataset_code": dataset_code,
                "dataset": dataset,
                "scenario_code": scenario_code,
                "scenario": scenario,
                "best_method_by_f1_auc_acc": best["method"],
                "exponential_rank_by_f1_auc_acc": exp_rank,
                "exponential_is_best_by_f1": bool(exp["f1"] == group["f1"].max()),
                "exponential_f1": exp.get("f1", np.nan),
                "best_f1": group["f1"].max(),
                "exponential_auc": exp.get("auc", np.nan),
                "best_auc": group["auc"].max() if "auc" in group else np.nan,
                "exponential_precision": exp.get("precision", np.nan),
                "exponential_recall": exp.get("recall", np.nan),
                "dm_wins_vs_3_methods": dm_wins,
                "dm_significant_wins_p_lt_0_05": significant_wins,
                "dm_significant_losses_p_lt_0_05": significant_losses,
                "conclusion": (
                    "exponential outperforms by F1"
                    if bool(exp["f1"] == group["f1"].max())
                    else f"not F1-best; best is {best['method']}"
                ),
            }
        )

    detail = pd.DataFrame(detail_rows)
    if detail.empty:
        return detail, detail

    by_crash = (
        detail.groupby(["scenario_code", "scenario"], dropna=False)
        .agg(
            markets_tested=("dataset_code", "count"),
            exponential_best_f1_count=("exponential_is_best_by_f1", "sum"),
            mean_exponential_f1=("exponential_f1", "mean"),
            mean_best_f1=("best_f1", "mean"),
            dm_wins_total=("dm_wins_vs_3_methods", "sum"),
            dm_significant_wins_total=("dm_significant_wins_p_lt_0_05", "sum"),
            dm_significant_losses_total=("dm_significant_losses_p_lt_0_05", "sum"),
        )
        .reset_index()
    )
    by_crash["exponential_best_f1_rate"] = by_crash["exponential_best_f1_count"] / by_crash["markets_tested"]
    by_crash["summary"] = np.where(
        by_crash["exponential_best_f1_count"] == by_crash["markets_tested"],
        "exponential outperforms in all markets by F1",
        np.where(
            by_crash["exponential_best_f1_count"] > 0,
            "exponential outperforms in some markets by F1",
            "exponential does not outperform by F1 in this crash window",
        ),
    )
    return detail, by_crash


def read_job_results(config: dict) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    output_dir = BASE / config["stock_name"] / "out_put_four_gaf_fanogan"
    perf_path = output_dir / "comparison_summary.csv"
    if perf_path.exists():
        perf = pd.read_csv(perf_path)
    else:
        perf = pd.concat([pd.read_csv(output_dir / m / "metrics_summary.csv") for m in METHODS], ignore_index=True)

    score_frames = []
    for method in METHODS:
        score_path = output_dir / method / "test_scores.csv"
        score_frames.append(pd.read_csv(score_path))
    scores = pd.concat(score_frames, ignore_index=True)
    dm = dm_tests(scores)

    for table in (perf, dm, scores):
        table.insert(0, "scenario", config["scenario"])
        table.insert(0, "scenario_code", config["scenario_code"])
        table.insert(0, "dataset", config["dataset"])
        table.insert(0, "dataset_code", config["dataset_code"])
    return perf, dm, scores


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center")
        for col_idx, cells in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in cells)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 56)
    wb.save(path)


def save_job_report(config: dict, perf: pd.DataFrame, dm: pd.DataFrame, scores: pd.DataFrame) -> Path:
    report = BASE / config["stock_name"] / f"{config['stock_name']}_performance_dm_test.xlsx"
    run_config = pd.DataFrame([config])
    with pd.ExcelWriter(report, engine="openpyxl") as writer:
        run_config.to_excel(writer, sheet_name="run_config", index=False)
        perf.to_excel(writer, sheet_name="performance", index=False)
        dm.to_excel(writer, sheet_name="dm_test", index=False)
        scores.to_excel(writer, sheet_name="test_scores_all_methods", index=False)
    style_workbook(report)
    return report


def outputs_complete(stock_name: str) -> bool:
    output_dir = BASE / stock_name / "out_put_four_gaf_fanogan"
    return all((output_dir / method / "metrics_summary.csv").exists() and (output_dir / method / "test_scores.csv").exists() for method in METHODS)


def run_job(config: dict, force: bool) -> None:
    stock_dir = BASE / config["stock_name"]
    if outputs_complete(config["stock_name"]) and not force:
        print(f"[SKIP] Existing complete output: {config['stock_name']}", flush=True)
        return

    if stock_dir.exists() and not outputs_complete(config["stock_name"]):
        shutil.rmtree(stock_dir)
    elif stock_dir.exists() and force:
        shutil.rmtree(stock_dir)

    env = os.environ.copy()
    env.update(
        {
            "EXCEL_SHEET": "0",
            "EXCEL_FIELD": "closed",
            "DATE_COLUMN": "date",
            "TRAIN_END_DATE": config["train_end_table"],
            "TEST_START_DATE": config["test_start_table"],
            "FORECAST_DAYS": str(config["forecast_days"]),
            "CRASH_START": str(config["crash_start_index"]),
            "CRASH_END": str(config["crash_end_index_exclusive"]),
            "CRASH_ONSET_DATE": str(config.get("crash_onset_date", "")),
            "CRASH_END_DATE": str(config.get("crash_end_date", "")),
            "PYTHONUNBUFFERED": "1",
        }
    )

    print(
        f"\n[JOB] {config['dataset']} | {config['scenario']} | "
        f"test={config['actual_test_start']} to {config['actual_test_end']} | "
        f"windows={config['forecast_days']} crash={config['n_test_crash_windows']} normal={config['n_test_normal_windows']}",
        flush=True,
    )
    subprocess.run(
        [
            sys.executable,
            str(SINGLE_JOB),
            "--input",
            config["input"],
            "--stock-name",
            config["stock_name"],
        ],
        cwd=str(BASE),
        env=env,
        check=True,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Run f-AnoGAN for 3 datasets across 3 market-crash windows.")
    parser.add_argument("--force", action="store_true", help="Delete and rerun existing completed outputs.")
    parser.add_argument("--summary-only", action="store_true", help="Do not train; only rebuild Excel reports from completed outputs.")
    parser.add_argument("--wait-existing", action="store_true", help="Wait for other f-AnoGAN jobs before starting this batch.")
    parser.add_argument("--dry-run", action="store_true", help="Only print the job configuration; do not train or build reports.")
    args = parser.parse_args()

    os.chdir(BASE)
    configs = [compute_job_config(dataset, scenario) for dataset in DATASETS for scenario in SCENARIOS]

    config_table = pd.DataFrame(configs)
    print("\n[CONFIG] Jobs to run")
    print(
        config_table[
            [
                "dataset",
                "scenario",
                "actual_train_start",
                "actual_train_end",
                "actual_test_start",
                "actual_test_end",
                "forecast_days",
                "n_train_normal_windows",
                "n_test_crash_windows",
                "n_test_normal_windows",
            ]
        ].to_string(index=False),
        flush=True,
    )

    if args.dry_run:
        print("\n[DRY-RUN] No training started.")
        return

    if not args.summary_only:
        if args.wait_existing:
            wait_for_external_fanogan_runs()
        for config in configs:
            run_job(config, args.force)

    all_perf = []
    all_dm = []
    all_reports = []
    for config in configs:
        perf, dm, scores = read_job_results(config)
        all_perf.append(perf)
        all_dm.append(dm)
        all_reports.append({"stock_name": config["stock_name"], "report": str(save_job_report(config, perf, dm, scores))})

    perf_all = pd.concat(all_perf, ignore_index=True)
    dm_all = pd.concat(all_dm, ignore_index=True)
    exp_detail, exp_by_crash = exponential_outperformance_summary(perf_all, dm_all)
    reports = pd.DataFrame(all_reports)

    with pd.ExcelWriter(MASTER_REPORT, engine="openpyxl") as writer:
        config_table.to_excel(writer, sheet_name="run_config", index=False)
        perf_all.to_excel(writer, sheet_name="all_performance", index=False)
        dm_all.to_excel(writer, sheet_name="all_dm_test", index=False)
        exp_detail.to_excel(writer, sheet_name="exp_outperform_detail", index=False)
        exp_by_crash.to_excel(writer, sheet_name="exp_outperform_by_crash", index=False)
        reports.to_excel(writer, sheet_name="job_reports", index=False)
        for scenario in SCENARIOS:
            sc = scenario["scenario"]
            perf_all.loc[perf_all["scenario"] == sc].to_excel(writer, sheet_name=f"{scenario['code'][:20]}_perf", index=False)
            dm_all.loc[dm_all["scenario"] == sc].to_excel(writer, sheet_name=f"{scenario['code'][:20]}_dm", index=False)
    style_workbook(MASTER_REPORT)

    print("\n[MASTER] Performance summary")
    show_cols = ["dataset", "scenario", "method", "accuracy", "precision", "recall", "f1", "auc"]
    print(perf_all[[c for c in show_cols if c in perf_all.columns]].to_string(index=False))
    print("\n[MASTER] DM test summary")
    print(dm_all.to_string(index=False))
    print("\n[MASTER] Exponential outperformance by crash")
    print(exp_by_crash.to_string(index=False))
    print("\n[MASTER] Exponential outperformance detail")
    print(exp_detail.to_string(index=False))
    print(f"\n[MASTER] Excel saved: {MASTER_REPORT}")


if __name__ == "__main__":
    main()
