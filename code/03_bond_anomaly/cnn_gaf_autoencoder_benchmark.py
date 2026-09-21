from __future__ import annotations

import argparse
import math
import os
import random
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.metrics import accuracy_score, f1_score, precision_score, recall_score, roc_auc_score
from torch.utils.data import DataLoader, TensorDataset


BASE = Path(__file__).resolve().parent
# --- replication package (added): data in <repo>/input, results in <repo>/output/03_bond_anomaly ---
import os as _rp_os, sys as _rp_sys
from pathlib import Path as _RpPath
REPO_ROOT = _RpPath(__file__).resolve().parents[2]
REPO_INPUT = REPO_ROOT / "input"
REPO_OUTPUT = REPO_ROOT / "output" / "03_bond_anomaly"
REPO_OUTPUT.mkdir(parents=True, exist_ok=True)
_rp_sys.path.insert(0, str(REPO_ROOT / "code" / "common"))
# ------------------------------------------------------------------------------------------
OUTPUT_ROOT = REPO_OUTPUT / "cnn_benchmark_outputs"
METHODS = ["cosine", "exponential", "arctan", "arccosh"]
WINDOW_SIZE = 32
IMG_SIZE = 32


@dataclass
class JobConfig:
    job_name: str
    dataset: str
    scenario: str
    input_file: Path
    excel_sheet: str | int
    price_col: str
    date_col: str
    train_start: str
    train_end: str
    test_start: str
    test_end: str | None = None
    forecast_days: int | None = None
    crash_start_index: int | None = None
    crash_end_index_exclusive: int | None = None
    crash_start_date: str | None = None
    crash_end_date: str | None = None


class ConvAutoEncoder(nn.Module):
    def __init__(self) -> None:
        super().__init__()
        self.encoder = nn.Sequential(
            nn.Conv2d(1, 8, kernel_size=3, stride=2, padding=1),
            nn.ReLU(inplace=True),
            nn.Conv2d(8, 16, kernel_size=3, stride=2, padding=1),
            nn.ReLU(inplace=True),
            nn.Conv2d(16, 32, kernel_size=3, stride=2, padding=1),
            nn.ReLU(inplace=True),
        )
        self.decoder = nn.Sequential(
            nn.ConvTranspose2d(32, 16, kernel_size=4, stride=2, padding=1),
            nn.ReLU(inplace=True),
            nn.ConvTranspose2d(16, 8, kernel_size=4, stride=2, padding=1),
            nn.ReLU(inplace=True),
            nn.ConvTranspose2d(8, 1, kernel_size=4, stride=2, padding=1),
            nn.Tanh(),
        )

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        return self.decoder(self.encoder(x))


def set_seed(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)


def parse_date_series(series: pd.Series) -> pd.Series:
    text = series.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    parsed = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")

    be_mask = text.str.fullmatch(r"(24|25|26)\d{6}")
    if be_mask.any():
        be_text = text[be_mask]
        be_year = pd.to_numeric(be_text.str.slice(0, 4), errors="coerce") - 543
        gregorian = (
            be_year.astype("Int64").astype(str)
            + "-"
            + be_text.str.slice(4, 6)
            + "-"
            + be_text.str.slice(6, 8)
        )
        parsed.loc[be_mask] = pd.to_datetime(gregorian, format="%Y-%m-%d", errors="coerce")

    for formatted_text, fmt in [
        (text.str.zfill(6), "%d%m%y"),
        (text.str.zfill(8), "%d%m%Y"),
        (text.str.zfill(8), "%Y%m%d"),
    ]:
        parsed = parsed.fillna(pd.to_datetime(formatted_text, format=fmt, errors="coerce"))
    return parsed.fillna(pd.to_datetime(text, errors="coerce", dayfirst=True))


def minmax_01(x: np.ndarray) -> np.ndarray:
    lo = float(np.min(x))
    hi = float(np.max(x))
    if hi - lo < 1e-12:
        return np.zeros_like(x, dtype=np.float32)
    return ((x - lo) / (hi - lo)).astype(np.float32)


def resize_1d(x: np.ndarray, out_size: int) -> np.ndarray:
    if len(x) == out_size:
        return x.astype(np.float32)
    src = np.linspace(0.0, 1.0, len(x))
    dst = np.linspace(0.0, 1.0, out_size)
    return np.interp(dst, src, x).astype(np.float32)


def gaf_cosine(series: np.ndarray, out_size: int = IMG_SIZE) -> np.ndarray:
    s = np.clip(2.0 * minmax_01(series) - 1.0, -1.0, 1.0)
    phi = np.arccos(s)
    phi = resize_1d(phi, out_size)
    return np.cos(phi[:, None] + phi[None, :]).astype(np.float32)


def gaf_arctan(series: np.ndarray, out_size: int = IMG_SIZE) -> np.ndarray:
    phi = resize_1d(np.arctan(minmax_01(series)), out_size)
    return np.cos(phi[:, None] + phi[None, :]).astype(np.float32)


def gaf_arccosh(series: np.ndarray, out_size: int = IMG_SIZE) -> np.ndarray:
    phi = resize_1d(np.arccosh(1.0 + minmax_01(series)), out_size)
    return np.cos(phi[:, None] + phi[None, :]).astype(np.float32)


def gaf_exponential(series: np.ndarray, out_size: int = IMG_SIZE) -> np.ndarray:
    phi = np.pi * (np.exp(minmax_01(series)) - 1.0) / (np.e - 1.0)
    phi = resize_1d(phi, out_size)
    return np.cos(phi[:, None] + phi[None, :]).astype(np.float32)


GAF_METHODS: dict[str, Callable[[np.ndarray, int], np.ndarray]] = {
    "cosine": gaf_cosine,
    "exponential": gaf_exponential,
    "arctan": gaf_arctan,
    "arccosh": gaf_arccosh,
}


def load_prices_and_dates(config: JobConfig) -> tuple[np.ndarray, np.ndarray]:
    sheet = config.excel_sheet
    df = pd.read_excel(config.input_file, sheet_name=sheet)
    cols = {str(c).strip().lower(): c for c in df.columns}
    price_col = cols[config.price_col.lower()]
    date_col = cols[config.date_col.lower()]
    prices = pd.to_numeric(df[price_col], errors="coerce")
    dates = parse_date_series(df[date_col])
    valid = prices.notna() & dates.notna()
    out = pd.DataFrame({"Date": dates[valid], "price": prices[valid].astype(float)})
    out = out.sort_values("Date").drop_duplicates("Date", keep="last").reset_index(drop=True)
    return out["price"].to_numpy(dtype=np.float64), out["Date"].to_numpy()


def make_split(config: JobConfig, prices: np.ndarray, x_axis: np.ndarray) -> dict[str, np.ndarray]:
    starts = np.arange(len(prices) - WINDOW_SIZE + 1, dtype=np.int64)
    start_dates = pd.to_datetime(x_axis[starts])
    train_end = pd.Timestamp(config.train_end)
    test_start = pd.Timestamp(config.test_start)

    if config.forecast_days is not None:
        candidates = np.where(start_dates >= test_start)[0]
        test_idx = candidates[: config.forecast_days]
    else:
        test_end = pd.Timestamp(config.test_end)
        test_idx = np.where((start_dates >= test_start) & (start_dates <= test_end))[0]

    if len(test_idx) == 0:
        raise RuntimeError(f"No test windows found for {config.job_name}")

    if config.crash_start_index is not None and config.crash_end_index_exclusive is not None:
        crash_start = int(config.crash_start_index)
        crash_end = int(config.crash_end_index_exclusive)
    else:
        crash_start_date = pd.Timestamp(config.crash_start_date or config.test_start)
        crash_end_date = pd.Timestamp(config.crash_end_date or config.test_end)
        all_dates = pd.to_datetime(x_axis)
        crash_start = int(np.where(all_dates >= crash_start_date)[0][0])
        after = np.where(all_dates > crash_end_date)[0]
        crash_end = int(after[0]) if len(after) else int(len(all_dates))

    labels = np.array([1 if ((s < crash_end) and (s + WINDOW_SIZE > crash_start)) else 0 for s in starts], dtype=np.int64)
    train_mask = (start_dates <= train_end) & (labels == 0)
    test_mask = np.zeros(len(starts), dtype=bool)
    test_mask[test_idx] = True
    return {
        "starts": starts,
        "labels": labels,
        "train_mask": train_mask,
        "test_mask": test_mask,
        "start_dates": start_dates.to_numpy(),
        "crash_start": np.array([crash_start]),
        "crash_end": np.array([crash_end]),
    }


def build_gafs(prices: np.ndarray, starts: np.ndarray, gaf_fn: Callable[[np.ndarray, int], np.ndarray]) -> np.ndarray:
    return np.stack([gaf_fn(prices[s : s + WINDOW_SIZE], IMG_SIZE) for s in starts]).astype(np.float32)


def score_autoencoder(model: nn.Module, x: np.ndarray, device: torch.device, batch_size: int) -> np.ndarray:
    loader = DataLoader(TensorDataset(torch.tensor(x).unsqueeze(1)), batch_size=batch_size, shuffle=False)
    scores: list[np.ndarray] = []
    model.eval()
    with torch.no_grad():
        for (xb,) in loader:
            xb = xb.to(device)
            recon = model(xb)
            loss = ((recon - xb) ** 2).mean(dim=(1, 2, 3))
            scores.append(loss.detach().cpu().numpy())
    return np.concatenate(scores)


def run_method(
    config: JobConfig,
    method: str,
    prices: np.ndarray,
    split: dict[str, np.ndarray],
    device: torch.device,
    epochs: int,
    batch_size: int,
    threshold_q: float,
) -> tuple[dict, pd.DataFrame]:
    gafs = build_gafs(prices, split["starts"], GAF_METHODS[method])
    x_train = gafs[split["train_mask"]]
    x_test = gafs[split["test_mask"]]
    y_test = split["labels"][split["test_mask"]]
    starts_test = split["starts"][split["test_mask"]]

    model = ConvAutoEncoder().to(device)
    loader = DataLoader(TensorDataset(torch.tensor(x_train).unsqueeze(1)), batch_size=batch_size, shuffle=True)
    optim = torch.optim.Adam(model.parameters(), lr=1e-3)
    loss_fn = nn.MSELoss()
    model.train()
    for _ in range(epochs):
        for (xb,) in loader:
            xb = xb.to(device)
            optim.zero_grad(set_to_none=True)
            loss = loss_fn(model(xb), xb)
            loss.backward()
            optim.step()

    train_scores = score_autoencoder(model, x_train, device, batch_size)
    test_scores = score_autoencoder(model, x_test, device, batch_size)
    threshold = float(np.quantile(train_scores, threshold_q))
    y_pred = (test_scores > threshold).astype(np.int64)

    try:
        auc = float(roc_auc_score(y_test, test_scores))
    except Exception:
        auc = float("nan")

    metrics = {
        "job_name": config.job_name,
        "dataset": config.dataset,
        "scenario": config.scenario,
        "model": "cnn_autoencoder",
        "method": method,
        "n_train_normal": int(len(x_train)),
        "n_test_total": int(len(x_test)),
        "n_test_normal": int((y_test == 0).sum()),
        "n_test_crash": int((y_test == 1).sum()),
        "threshold": threshold,
        "accuracy": float(accuracy_score(y_test, y_pred)),
        "precision": float(precision_score(y_test, y_pred, zero_division=0)),
        "recall": float(recall_score(y_test, y_pred, zero_division=0)),
        "f1": float(f1_score(y_test, y_pred, zero_division=0)),
        "auc": auc,
        "mean_train_score": float(np.mean(train_scores)),
        "mean_test_score": float(np.mean(test_scores)),
    }
    details = pd.DataFrame(
        {
            "job_name": config.job_name,
            "dataset": config.dataset,
            "scenario": config.scenario,
            "model": "cnn_autoencoder",
            "method": method,
            "window_start": starts_test,
            "window_end": starts_test + WINDOW_SIZE - 1,
            "label_0normal_1crash": y_test,
            "anomaly_score": test_scores,
            "threshold": threshold,
            "predicted_label": y_pred,
        }
    )
    return metrics, details


def normal_cdf(x: float) -> float:
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def dm_tests(scores: pd.DataFrame) -> pd.DataFrame:
    wide = None
    for method in METHODS:
        df = scores.loc[scores["method"] == method, ["window_start", "window_end", "label_0normal_1crash", "predicted_label"]].copy()
        df[f"loss_{method}"] = (df["label_0normal_1crash"] - df["predicted_label"]) ** 2
        df = df[["window_start", "window_end", f"loss_{method}"]]
        wide = df if wide is None else wide.merge(df, on=["window_start", "window_end"], how="inner")

    rows = []
    for other in ["cosine", "arctan", "arccosh"]:
        d = wide[f"loss_{other}"] - wide["loss_exponential"]
        t = int(len(d))
        mean_diff = float(d.mean())
        var = float(d.var(ddof=1)) if t > 1 else float("nan")
        if t > 1 and var > 0:
            dm_stat = mean_diff / math.sqrt(var / t)
            p_value = 2.0 * (1.0 - normal_cdf(abs(dm_stat)))
        elif mean_diff == 0:
            dm_stat = 0.0
            p_value = 1.0
        else:
            dm_stat = math.copysign(float("inf"), mean_diff)
            p_value = 0.0
        exp_loss = float(wide["loss_exponential"].mean())
        other_loss = float(wide[f"loss_{other}"].mean())
        rows.append(
            {
                "comparison": f"exponential_vs_{other}",
                "T": t,
                "loss_exponential": exp_loss,
                "loss_other": other_loss,
                "mean_loss_diff_other_minus_exp": mean_diff,
                "dm_statistic": dm_stat,
                "p_value": p_value,
                "winner_lower_loss": "exponential" if exp_loss < other_loss else other if other_loss < exp_loss else "tie",
                "interpretation": "positive DM favors exponential" if mean_diff > 0 else "negative DM favors other" if mean_diff < 0 else "same loss",
            }
        )
    return pd.DataFrame(rows)


def exponential_summary(perf: pd.DataFrame, dm: pd.DataFrame) -> pd.DataFrame:
    ranked = perf.sort_values(["f1", "auc", "accuracy"], ascending=False).reset_index(drop=True)
    exp = ranked.loc[ranked["method"] == "exponential"].iloc[0]
    best = ranked.iloc[0]
    return pd.DataFrame(
        [
            {
                "job_name": exp["job_name"],
                "dataset": exp["dataset"],
                "scenario": exp["scenario"],
                "model": exp["model"],
                "best_method_by_f1_auc_acc": best["method"],
                "exponential_rank_by_f1_auc_acc": int(ranked.index[ranked["method"] == "exponential"][0]) + 1,
                "exponential_is_best_by_f1": bool(exp["f1"] == perf["f1"].max()),
                "exponential_f1": exp["f1"],
                "best_f1": perf["f1"].max(),
                "exponential_auc": exp["auc"],
                "best_auc": perf["auc"].max(),
                "dm_wins_vs_3_methods": int((dm["winner_lower_loss"] == "exponential").sum()),
                "dm_significant_wins_p_lt_0_05": int(((dm["winner_lower_loss"] == "exponential") & (dm["p_value"] < 0.05)).sum()),
                "conclusion": "exponential outperforms by F1" if bool(exp["f1"] == perf["f1"].max()) else f"not F1-best; best is {best['method']}",
            }
        ]
    )


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center")
        for col_idx, cells in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in cells)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 55)
    wb.save(path)


def us1m_job() -> JobConfig:
    return JobConfig(
        job_name="US1M_COVID200_CNN_AE",
        dataset="US1M",
        scenario="COVID-19 market crash",
        input_file=REPO_INPUT / "bond_prepared" / "US1M_COVID200_daily_final2.xlsx",
        excel_sheet=0,
        price_col="cp",
        date_col="Date",
        train_start="2015-05-06",
        train_end="2019-12-30",
        test_start="2020-01-02",
        forecast_days=200,
        crash_start_index=938,
        crash_end_index_exclusive=959,
    )


def market3_jobs() -> list[JobConfig]:
    datasets = [
        ("S&P500", "SP500_U500", REPO_INPUT / "u500_1day.xlsx"),
        ("DAX INDEX", "DAX_INDEX", REPO_INPUT / "dax_1day.xlsx"),
        ("DJIA", "DJIA_U30", REPO_INPUT / "u30_1day.xlsx"),
    ]
    scenarios = [
        ("COVID-19 market crash", "COVID19_market_crash", "2013-11-01", "2019-12-05", "2019-12-06", "2020-03-31"),
        ("Russia-Ukraine war crash", "Russia_Ukraine_war_crash", "2013-11-01", "2018-02-10", "2021-10-28", "2022-02-24"),
        ("Chinese real asset market crash", "Chinese_real_asset_market_crash", "2013-11-01", "2018-02-10", "2023-07-05", "2023-10-31"),
    ]
    jobs: list[JobConfig] = []
    for dataset, dataset_code, path in datasets:
        for scenario, scenario_code, train_start, train_end, test_start, test_end in scenarios:
            jobs.append(
                JobConfig(
                    job_name=f"{dataset_code}_{scenario_code}_CNN_AE",
                    dataset=dataset,
                    scenario=scenario,
                    input_file=path,
                    excel_sheet=0,
                    price_col="closed",
                    date_col="date",
                    train_start=train_start,
                    train_end=train_end,
                    test_start=test_start,
                    test_end=test_end,
                    crash_start_date=test_start,
                    crash_end_date=test_end,
                )
            )
    return jobs


def run_job(config: JobConfig, device: torch.device, epochs: int, batch_size: int, threshold_q: float) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    prices, x_axis = load_prices_and_dates(config)
    split = make_split(config, prices, x_axis)
    metrics_rows = []
    score_frames = []
    for method in METHODS:
        print(f"[CNN] {config.job_name} | {method}", flush=True)
        metrics, details = run_method(config, method, prices, split, device, epochs, batch_size, threshold_q)
        metrics_rows.append(metrics)
        score_frames.append(details)
    perf = pd.DataFrame(metrics_rows).sort_values(["f1", "auc", "accuracy"], ascending=False).reset_index(drop=True)
    scores = pd.concat(score_frames, ignore_index=True)
    dm = dm_tests(scores)
    summary = exponential_summary(perf, dm)

    out_dir = OUTPUT_ROOT / config.job_name
    out_dir.mkdir(parents=True, exist_ok=True)
    perf.to_csv(out_dir / "cnn_performance.csv", index=False)
    dm.to_csv(out_dir / "cnn_dm_test.csv", index=False)
    scores.to_csv(out_dir / "cnn_test_scores.csv", index=False)
    summary.to_csv(out_dir / "cnn_exponential_summary.csv", index=False)
    report = out_dir / f"{config.job_name}_cnn_benchmark.xlsx"
    with pd.ExcelWriter(report, engine="openpyxl") as writer:
        pd.DataFrame([asdict(config)]).to_excel(writer, sheet_name="run_config", index=False)
        perf.to_excel(writer, sheet_name="performance", index=False)
        dm.to_excel(writer, sheet_name="dm_test", index=False)
        summary.to_excel(writer, sheet_name="exponential_summary", index=False)
        scores.to_excel(writer, sheet_name="test_scores_all_methods", index=False)
    style_workbook(report)
    return perf, dm, summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Fast CNN autoencoder benchmark for GAF mappings.")
    parser.add_argument("--preset", choices=["us1m", "market3"], default="us1m")
    parser.add_argument("--epochs", type=int, default=30)
    parser.add_argument("--batch-size", type=int, default=64)
    parser.add_argument("--threshold-q", type=float, default=0.95)
    parser.add_argument("--device", choices=["auto", "cuda", "cpu"], default="auto")
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()

    os.chdir(BASE)
    set_seed(args.seed)
    if args.device == "cuda" or (args.device == "auto" and torch.cuda.is_available()):
        device = torch.device("cuda")
    else:
        device = torch.device("cpu")
    print(f"[CNN] Device: {device} | epochs={args.epochs} | preset={args.preset}", flush=True)

    jobs = [us1m_job()] if args.preset == "us1m" else market3_jobs()
    all_perf = []
    all_dm = []
    all_summary = []
    for job in jobs:
        perf, dm, summary = run_job(job, device, args.epochs, args.batch_size, args.threshold_q)
        all_perf.append(perf)
        all_dm.append(dm.assign(job_name=job.job_name, dataset=job.dataset, scenario=job.scenario))
        all_summary.append(summary)

    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    master = OUTPUT_ROOT / f"cnn_benchmark_{args.preset}_summary.xlsx"
    perf_all = pd.concat(all_perf, ignore_index=True)
    dm_all = pd.concat(all_dm, ignore_index=True)
    summary_all = pd.concat(all_summary, ignore_index=True)
    with pd.ExcelWriter(master, engine="openpyxl") as writer:
        perf_all.to_excel(writer, sheet_name="all_performance", index=False)
        dm_all.to_excel(writer, sheet_name="all_dm_test", index=False)
        summary_all.to_excel(writer, sheet_name="exponential_summary", index=False)
    style_workbook(master)

    print("\n[CNN] Performance")
    print(perf_all.to_string(index=False))
    print("\n[CNN] DM test")
    print(dm_all.to_string(index=False))
    print("\n[CNN] Exponential summary")
    print(summary_all.to_string(index=False))
    print(f"\n[CNN] Excel saved: {master}", flush=True)


if __name__ == "__main__":
    main()
