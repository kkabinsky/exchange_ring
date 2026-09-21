from __future__ import annotations

import argparse
import importlib
import math
import os
import sys
import time
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except Exception:  # pragma: no cover
    load_workbook = None

BASE = Path(__file__).resolve().parent
# --- replication package (added): data in <repo>/input, results in <repo>/output/03_bond_anomaly ---
import os as _rp_os, sys as _rp_sys
from pathlib import Path as _RpPath
REPO_ROOT = _RpPath(__file__).resolve().parents[2]
REPO_INPUT = REPO_ROOT / "input"
REPO_OUTPUT = REPO_ROOT / "output" / "03_bond_anomaly"
REPO_OUTPUT.mkdir(parents=True, exist_ok=True)
_rp_sys.path.insert(0, str(REPO_ROOT / "code" / "common"))
# ------------------------------------------------------------------------------------------
COMBINE_TTM = REPO_INPUT / "combine_TTM.xlsx"
OUT_ROOT = REPO_OUTPUT / "bond_1m_30y_compare_e20_daily"
INPUT_DIR = OUT_ROOT / "prepared_inputs"  # written by --stage prepare
HYBRID_DIR = REPO_ROOT / "code" / "04_cross_market"  # hybrid pipeline code
HYBRID_INPUT_DIR = OUT_ROOT / "prepared_inputs_hybrid"
CNN_ROOT = OUT_ROOT / "cnn_autoencoder"
FANOGAN_ROOT = OUT_ROOT / "standalone_fanogan"
HYBRID_ROOT = OUT_ROOT / "hybrid"
SUMMARY_XLSX = OUT_ROOT / "bond_1m_30y_cnn_fanogan_hybrid_summary_latest.xlsx"
SUMMARY_CSV_PREFIX = OUT_ROOT / "bond_1m_30y_summary"

METHODS = ["cosine", "exponential", "arctan", "arccosh"]
DATASETS = [
    {
        "dataset": "US1M",
        "asset": "US Treasury yield 1 month",
        "column": "1M",
        "stock": "US1M_BOND_COVID200_daily",
    },
    {
        "dataset": "US30Y",
        "asset": "US Treasury yield 30 year",
        "column": "30Y",
        "stock": "US30Y_BOND_COVID200_daily",
    },
]

START_DATE = "2015-05-06"
END_DATE = "2020-12-14"
TRAIN_END_DATE = "2019-12-30"
TEST_START_DATE = "2020-01-02"
FORECAST_DAYS = 200
CRASH_START_DATE = "2020-03-02"
CRASH_END_DATE = "2020-03-30"
WINDOW_SIZE = 32
EPOCHS_SMALL = 20
CNN_EPOCHS = 30


def log(msg: str) -> None:
    print(msg, flush=True)


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def method_norm(x: object) -> str:
    return str(x).strip().lower()


def first_existing_column(df: pd.DataFrame, names: Iterable[str]) -> Optional[str]:
    lower_map = {str(c).lower(): c for c in df.columns}
    for name in names:
        if name in df.columns:
            return name
        if str(name).lower() in lower_map:
            return lower_map[str(name).lower()]
    return None


def normalize_dates(series: pd.Series) -> pd.Series:
    if np.issubdtype(series.dtype, np.number):
        return pd.to_datetime(series.astype("Int64").astype(str), errors="coerce")
    return pd.to_datetime(series, errors="coerce")


def prepare_inputs() -> pd.DataFrame:
    ensure_dir(OUT_ROOT)
    ensure_dir(INPUT_DIR)
    ensure_dir(HYBRID_INPUT_DIR)
    if not COMBINE_TTM.exists():
        raise FileNotFoundError(f"combine_TTM not found: {COMBINE_TTM}")

    raw = pd.read_excel(COMBINE_TTM, sheet_name="Ycurve")
    if "Date" not in raw.columns:
        raise KeyError("combine_TTM.xlsx sheet Ycurve must contain Date column")
    raw = raw.copy()
    raw["Date"] = normalize_dates(raw["Date"])
    raw = raw.dropna(subset=["Date"]).sort_values("Date").reset_index(drop=True)

    start = pd.Timestamp(START_DATE)
    end = pd.Timestamp(END_DATE)
    train_end = pd.Timestamp(TRAIN_END_DATE)
    test_start = pd.Timestamp(TEST_START_DATE)
    crash_start = pd.Timestamp(CRASH_START_DATE)
    crash_end = pd.Timestamp(CRASH_END_DATE)

    rows = []
    for ds in DATASETS:
        col = ds["column"]
        if col not in raw.columns:
            raise KeyError(f"Column {col} not found in combine_TTM.xlsx")
        sub = raw.loc[(raw["Date"] >= start) & (raw["Date"] <= end), ["Date", col]].copy()
        sub[col] = pd.to_numeric(sub[col], errors="coerce")
        sub = sub.dropna(subset=[col]).sort_values("Date").drop_duplicates("Date", keep="last").reset_index(drop=True)
        sub = sub.rename(columns={col: "cp"})
        date_s = pd.to_datetime(sub["Date"]).reset_index(drop=True)
        # The original f-AnoGAN loader parses compact YYYYMMDD reliably; hyphenated dates can be treated as day-first.
        sub["Date"] = date_s.dt.strftime("%Y%m%d").astype(int)

        out_file = INPUT_DIR / f"{ds['stock']}_final2.xlsx"
        hybrid_file = HYBRID_INPUT_DIR / f"{ds['stock']}_final2.xlsx"
        sub.to_excel(out_file, index=False)
        sub.to_excel(hybrid_file, index=False)

        test_start_idx_arr = np.flatnonzero(date_s >= test_start)
        crash_start_idx_arr = np.flatnonzero(date_s >= crash_start)
        crash_end_idx_arr = np.flatnonzero(date_s > crash_end)
        if len(test_start_idx_arr) == 0 or len(crash_start_idx_arr) == 0 or len(crash_end_idx_arr) == 0:
            raise ValueError(f"Cannot locate test/crash indices for {ds['dataset']}")
        test_start_idx = int(test_start_idx_arr[0])
        crash_start_idx = int(crash_start_idx_arr[0])
        crash_end_idx = int(crash_end_idx_arr[0])
        n_price = len(sub)
        n_windows = max(n_price - WINDOW_SIZE + 1, 0)
        candidate_starts = np.arange(n_windows)
        test_window_starts = candidate_starts[test_start_idx <= candidate_starts][:FORECAST_DAYS]
        labels = [1 if ((int(s) < crash_end_idx) and (int(s) + WINDOW_SIZE > crash_start_idx)) else 0 for s in test_window_starts]

        rows.append({
            "dataset": ds["dataset"],
            "asset": ds["asset"],
            "source_file": str(COMBINE_TTM),
            "source_sheet": "Ycurve",
            "source_column": col,
            "prepared_file": str(out_file),
            "hybrid_prepared_file": str(hybrid_file),
            "first_date": str(date_s.iloc[0].date()) if len(sub) else None,
            "last_date": str(date_s.iloc[-1].date()) if len(sub) else None,
            "n_prices": n_price,
            "n_windows": n_windows,
            "train_end_date": TRAIN_END_DATE,
            "test_start_date": TEST_START_DATE,
            "forecast_days": FORECAST_DAYS,
            "test_start_price_index": test_start_idx,
            "test_end_price_date": str(date_s.iloc[min(test_start_idx + FORECAST_DAYS - 1, len(date_s) - 1)].date()),
            "crash_start_date": CRASH_START_DATE,
            "crash_end_date": CRASH_END_DATE,
            "crash_start_price_index": crash_start_idx,
            "crash_end_price_index_exclusive": crash_end_idx,
            "test_windows": len(test_window_starts),
            "test_crash_windows": int(np.sum(labels)),
            "test_normal_windows": int(len(labels) - np.sum(labels)),
        })
        log(f"[PREP] {ds['dataset']}: wrote {out_file} | rows={n_price} | test={len(labels)} | crash_windows={int(np.sum(labels))}")

    meta = pd.DataFrame(rows)
    meta.to_csv(OUT_ROOT / "input_metadata.csv", index=False)
    return meta


def input_file_for(ds: Dict[str, str]) -> Path:
    return INPUT_DIR / f"{ds['stock']}_final2.xlsx"


def _set_common_fast_env(crash_start: int, crash_end: int) -> None:
    os.environ.update({
        "PYTHONUNBUFFERED": "1",
        "EXCEL_FIELD": "cp",
        "DATE_COLUMN": "Date",
        "TRAIN_END_DATE": TRAIN_END_DATE,
        "TEST_START_DATE": TEST_START_DATE,
        "FORECAST_DAYS": str(FORECAST_DAYS),
        "CRASH_START": str(crash_start),
        "CRASH_END": str(crash_end),
        "N_EPOCHS_GAN": str(EPOCHS_SMALL),
        "N_EPOCHS_ENC": str(EPOCHS_SMALL),
        "BATCH_SIZE": "32",
        "N_CRITIC": "1",
        "EXPORT_SAMPLE_GAF": "0",
        "SAVE_ALL_SAMPLE_IMAGES": "0",
        "QUICK": "0",
    })


def run_cnn(force: bool = False) -> None:
    meta = prepare_inputs()
    ensure_dir(CNN_ROOT)
    if str(BASE) not in sys.path:
        sys.path.insert(0, str(BASE))
    import torch
    import cnn_gaf_autoencoder_benchmark as cnn

    cnn.OUTPUT_ROOT = CNN_ROOT
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    log(f"[CNN] device={device} | output={CNN_ROOT}")

    for ds in DATASETS:
        mrow = meta.loc[meta["dataset"] == ds["dataset"]].iloc[0]
        out_dir = CNN_ROOT / f"{ds['stock']}_CNN_AE"
        done = out_dir / "cnn_performance.csv"
        if done.exists() and not force:
            log(f"[CNN] skip existing {ds['dataset']} -> {done}")
            continue
        cfg = cnn.JobConfig(
            job_name=f"{ds['stock']}_CNN_AE",
            dataset=ds["dataset"],
            scenario="COVID-19 market crash",
            input_file=str(input_file_for(ds)),
            excel_sheet=0,
            price_col="cp",
            date_col="Date",
            train_start=START_DATE,
            train_end=TRAIN_END_DATE,
            test_start=TEST_START_DATE,
            test_end=None,
            forecast_days=FORECAST_DAYS,
            crash_start_index=int(mrow["crash_start_price_index"]),
            crash_end_index_exclusive=int(mrow["crash_end_price_index_exclusive"]),
            crash_start_date=None,
            crash_end_date=None,
        )
        log(f"[CNN] run {ds['dataset']} -> {cfg.input_file}")
        cnn.run_job(cfg, device=device, epochs=CNN_EPOCHS, batch_size=32, threshold_q=0.95)


def run_fanogan(force: bool = False) -> None:
    meta = prepare_inputs()
    ensure_dir(FANOGAN_ROOT)
    if str(BASE) not in sys.path:
        sys.path.insert(0, str(BASE))
    os.chdir(FANOGAN_ROOT)

    first = meta.iloc[0]
    _set_common_fast_env(int(first["crash_start_price_index"]), int(first["crash_end_price_index_exclusive"]))
    if "fanogan_four_gaf_compare_2" in sys.modules:
        del sys.modules["fanogan_four_gaf_compare_2"]
    fanogan = importlib.import_module("fanogan_four_gaf_compare_2")
    log(f"[FANOGAN] device/module loaded | output={FANOGAN_ROOT}")

    for _, mrow in meta.iterrows():
        ds = next(x for x in DATASETS if x["dataset"] == mrow["dataset"])
        # The two prepared yield datasets use the same aligned indices. Keep this explicit in case source changes later.
        if int(mrow["crash_start_price_index"]) != int(first["crash_start_price_index"]):
            raise ValueError("Crash indices differ across datasets; run standalone f-AnoGAN per process to avoid imported env constants.")
        stock = ds["stock"]
        done = FANOGAN_ROOT / stock / "out_put_four_gaf_fanogan" / "comparison_summary.csv"
        if done.exists() and not force:
            log(f"[FANOGAN] skip existing {ds['dataset']} -> {done}")
            continue
        log(f"[FANOGAN] run {ds['dataset']} -> {input_file_for(ds)}")
        fanogan.run_one_stock(str(input_file_for(ds)), stock)


def run_hybrid(force: bool = False) -> None:
    prepare_inputs()
    ensure_dir(HYBRID_ROOT)
    reports = [HYBRID_ROOT / ds["stock"] / "hybrid_four_gaf_report.xlsx" for ds in DATASETS]
    if all(p.exists() for p in reports) and not force:
        log(f"[HYBRID] skip existing all reports -> {HYBRID_ROOT}")
        return

    os.environ.update({
        "PYTHONUNBUFFERED": "1",
        "EXCEL_DIR": str(HYBRID_INPUT_DIR),
        "BATCH_SUFFIX": "_final2.xlsx",
        "INPUT_EXCEL": "",
        "EXCEL_FIELD": "cp",
        "DATE_COLUMN": "Date",
        "OUT_ROOT": str(HYBRID_ROOT),
        "FORECAST_DAYS": str(FORECAST_DAYS),
        "BACKTEST_DAYS": str(FORECAST_DAYS),
        "MIN_BACKTEST_ANOMALIES": "0",
        "TADGAN_EPOCHS": str(EPOCHS_SMALL),
        "TADGAN_N_CRITIC": "1",
        "TADGAN_USE_CACHE": "0",
        "TADGAN_FORCE_TRAIN": "0",
        "N_EPOCHS_GAN": str(EPOCHS_SMALL),
        "N_EPOCHS_ENC": str(EPOCHS_SMALL),
        "N_CRITIC": "1",
        "FC_EPOCHS": "50",
        "BATCH_SIZE": "32",
        "FC_BATCH": "64",
        "SAVE_GAFS": "0",
        "QUICK": "0",
    })
    if str(HYBRID_DIR) not in sys.path:
        sys.path.insert(0, str(HYBRID_DIR))
    os.chdir(HYBRID_ROOT)
    if "hybrid_gaf_tadgan_fanogan" in sys.modules:
        del sys.modules["hybrid_gaf_tadgan_fanogan"]
    hybrid = importlib.import_module("hybrid_gaf_tadgan_fanogan")
    log(f"[HYBRID] batch run -> input={HYBRID_INPUT_DIR} | output={HYBRID_ROOT}")
    hybrid.main()


def auc_safe(y: np.ndarray, score: np.ndarray) -> float:
    try:
        from sklearn.metrics import roc_auc_score
        if len(np.unique(y)) < 2:
            return float("nan")
        return float(roc_auc_score(y, score))
    except Exception:
        return float("nan")


def rank01(values: Iterable[float]) -> np.ndarray:
    s = pd.Series(list(values), dtype="float64")
    return (s.rank(method="average").to_numpy(dtype=float) / (len(s) + 1.0))


def dm_exp_vs_baselines(score_tables: Dict[str, pd.DataFrame], label_col: str, score_col: str, pipeline: str, dataset: str) -> pd.DataFrame:
    rows = []
    if "exponential" not in score_tables:
        return pd.DataFrame(rows)
    n = min(len(df) for df in score_tables.values())
    if n <= 2:
        return pd.DataFrame(rows)
    exp_df = score_tables["exponential"].iloc[:n].copy()
    y = pd.to_numeric(exp_df[label_col], errors="coerce").to_numpy(dtype=float)
    losses = {}
    for method, df in score_tables.items():
        cur = df.iloc[:n].copy()
        score = pd.to_numeric(cur[score_col], errors="coerce").to_numpy(dtype=float)
        prob = rank01(score)
        losses[method] = (prob - y) ** 2
    exp_loss = losses["exponential"]
    for method in METHODS:
        if method == "exponential" or method not in losses:
            continue
        base_loss = losses[method]
        d = base_loss - exp_loss
        d = d[np.isfinite(d)]
        if len(d) <= 2:
            dm_stat = float("nan")
            p_one = float("nan")
        else:
            sd = float(np.std(d, ddof=1))
            if sd == 0 or not np.isfinite(sd):
                dm_stat = float("nan")
                p_one = float("nan")
            else:
                dm_stat = float(np.mean(d) / (sd / math.sqrt(len(d))))
                p_one = float(0.5 * math.erfc(dm_stat / math.sqrt(2.0)))
        rows.append({
            "pipeline": pipeline,
            "dataset": dataset,
            "comparison": f"exponential vs {method}",
            "baseline_method": method,
            "n": int(len(d)),
            "baseline_rank_brier_loss": float(np.nanmean(base_loss)),
            "exponential_rank_brier_loss": float(np.nanmean(exp_loss)),
            "loss_diff_baseline_minus_exp": float(np.nanmean(base_loss) - np.nanmean(exp_loss)),
            "dm_stat_positive_exp_better": dm_stat,
            "p_one_sided_exp_better": p_one,
            "exp_lower_loss": bool(np.nanmean(exp_loss) < np.nanmean(base_loss)),
            "significant_5pct": bool((p_one < 0.05) if np.isfinite(p_one) else False),
        })
    return pd.DataFrame(rows)


def tidy_perf_row(row: pd.Series, pipeline: str, dataset: str, asset: str, source_file: str) -> Dict[str, object]:
    def get(*names, default=np.nan):
        for name in names:
            if name in row.index:
                return row[name]
        return default
    method = method_norm(get("method", default=""))
    return {
        "pipeline": pipeline,
        "dataset": dataset,
        "asset": asset,
        "method": method,
        "accuracy": get("accuracy", "backtest_accuracy"),
        "precision": get("precision", "precision_crash", "backtest_precision_crash"),
        "recall": get("recall", "recall_crash", "backtest_recall_crash"),
        "f1": get("f1", "f1_crash", "backtest_f1_crash"),
        "auc": get("auc", "auc_crash", "backtest_auc_crash"),
        "tp": get("tp", "tp_crash", "backtest_tp_crash"),
        "fp": get("fp", "fp_crash", "backtest_fp_crash"),
        "tn": get("tn", "tn_normal", "backtest_tn_normal"),
        "fn": get("fn", "fn_crash", "backtest_fn_crash"),
        "n_test_total": get("n_test_total", "n_test", "backtest_n_samples"),
        "n_test_crash": get("n_test_crash", "backtest_n_crash"),
        "n_test_normal": get("n_test_normal", "backtest_n_normal"),
        "threshold": get("threshold"),
        "source_file": source_file,
    }


def collect_cnn(ds: Dict[str, str]) -> Tuple[pd.DataFrame, pd.DataFrame]:
    job = f"{ds['stock']}_CNN_AE"
    out_dir = CNN_ROOT / job
    perf_file = out_dir / "cnn_performance.csv"
    scores_file = out_dir / "cnn_test_scores.csv"
    perf = pd.DataFrame()
    dm = pd.DataFrame()
    if perf_file.exists():
        src = pd.read_csv(perf_file)
        perf = pd.DataFrame([tidy_perf_row(r, "CNN autoencoder", ds["dataset"], ds["asset"], str(perf_file)) for _, r in src.iterrows()])
    if scores_file.exists():
        scores = pd.read_csv(scores_file)
        method_col = first_existing_column(scores, ["method"])
        label_col = first_existing_column(scores, ["label_0normal_1crash", "true_label", "label"])
        score_col = first_existing_column(scores, ["anomaly_score", "score"])
        if method_col and label_col and score_col:
            tables = {}
            for method in METHODS:
                sub = scores.loc[scores[method_col].map(method_norm) == method].copy()
                if len(sub):
                    sub["y_crash"] = pd.to_numeric(sub[label_col], errors="coerce")
                    tables[method] = sub.rename(columns={score_col: "score"})[["y_crash", "score"]]
            dm = dm_exp_vs_baselines(tables, "y_crash", "score", "CNN autoencoder", ds["dataset"])
    return perf, dm


def collect_fanogan(ds: Dict[str, str]) -> Tuple[pd.DataFrame, pd.DataFrame]:
    out_dir = FANOGAN_ROOT / ds["stock"] / "out_put_four_gaf_fanogan"
    perf_file = out_dir / "comparison_summary.csv"
    perf = pd.DataFrame()
    dm = pd.DataFrame()
    if perf_file.exists():
        src = pd.read_csv(perf_file)
        perf = pd.DataFrame([tidy_perf_row(r, "f-AnoGAN standalone", ds["dataset"], ds["asset"], str(perf_file)) for _, r in src.iterrows()])
    tables = {}
    for method in METHODS:
        score_file = out_dir / method / "test_scores.csv"
        if score_file.exists():
            s = pd.read_csv(score_file)
            if "label_0normal_1crash" in s.columns and "anomaly_score" in s.columns:
                tables[method] = pd.DataFrame({
                    "y_crash": pd.to_numeric(s["label_0normal_1crash"], errors="coerce"),
                    "score": pd.to_numeric(s["anomaly_score"], errors="coerce"),
                })
    if tables:
        dm = dm_exp_vs_baselines(tables, "y_crash", "score", "f-AnoGAN standalone", ds["dataset"])
    return perf, dm


def collect_hybrid(ds: Dict[str, str]) -> Tuple[pd.DataFrame, pd.DataFrame]:
    report = HYBRID_ROOT / ds["stock"] / "hybrid_four_gaf_report.xlsx"
    perf = pd.DataFrame()
    dm = pd.DataFrame()
    if not report.exists():
        return perf, dm
    try:
        src = pd.read_excel(report, sheet_name="summary")
        perf = pd.DataFrame([tidy_perf_row(r, "Hybrid TadGAN-GAF-f-AnoGAN", ds["dataset"], ds["asset"], str(report)) for _, r in src.iterrows()])
    except Exception as e:
        log(f"[SUMMARY] Could not read hybrid summary {report}: {e}")

    tables = {}
    for method in METHODS:
        sheet = f"{method[:18]}_back"
        try:
            s = pd.read_excel(report, sheet_name=sheet)
        except Exception:
            continue
        if "true_label_user_0crash_1not" in s.columns and "fanogan_score" in s.columns:
            y = (pd.to_numeric(s["true_label_user_0crash_1not"], errors="coerce") == 0).astype(float)
            tables[method] = pd.DataFrame({"y_crash": y, "score": pd.to_numeric(s["fanogan_score"], errors="coerce")})
    if tables:
        dm = dm_exp_vs_baselines(tables, "y_crash", "score", "Hybrid TadGAN-GAF-f-AnoGAN", ds["dataset"])
    return perf, dm


def summarize() -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    ensure_dir(OUT_ROOT)
    meta_file = OUT_ROOT / "input_metadata.csv"
    meta = pd.read_csv(meta_file) if meta_file.exists() else prepare_inputs()
    meta_map = {r["dataset"]: r for _, r in meta.iterrows()}

    perf_frames: List[pd.DataFrame] = []
    dm_frames: List[pd.DataFrame] = []
    for ds in DATASETS:
        for collector in (collect_cnn, collect_fanogan, collect_hybrid):
            perf, dm = collector(ds)
            if len(perf):
                perf_frames.append(perf)
            if len(dm):
                dm_frames.append(dm)

    perf_all = pd.concat(perf_frames, ignore_index=True) if perf_frames else pd.DataFrame()
    dm_all = pd.concat(dm_frames, ignore_index=True) if dm_frames else pd.DataFrame()

    for col in ["accuracy", "precision", "recall", "f1", "auc", "threshold", "n_test_total", "n_test_crash", "n_test_normal"]:
        if col in perf_all.columns:
            perf_all[col] = pd.to_numeric(perf_all[col], errors="coerce")

    perf_main = perf_all.loc[perf_all.get("method", pd.Series(dtype=str)).isin(METHODS)].copy() if len(perf_all) else pd.DataFrame()
    if len(perf_main):
        perf_avg = perf_main.groupby(["pipeline", "method"], as_index=False).agg(
            accuracy_mean=("accuracy", "mean"),
            precision_mean=("precision", "mean"),
            recall_mean=("recall", "mean"),
            f1_mean=("f1", "mean"),
            auc_mean=("auc", "mean"),
            n_rows=("dataset", "count"),
        ).sort_values(["pipeline", "f1_mean", "auc_mean"], ascending=[True, False, False])
    else:
        perf_avg = pd.DataFrame()

    win_rows = []
    if len(perf_main):
        for (pipeline, dataset), sub in perf_main.groupby(["pipeline", "dataset"]):
            sub = sub.copy()
            exp = sub.loc[sub["method"] == "exponential"]
            best_f1 = sub.sort_values(["f1", "auc", "accuracy"], ascending=False).iloc[0]
            best_auc = sub.sort_values(["auc", "f1", "accuracy"], ascending=False).iloc[0]
            dm_sub = dm_all.loc[(dm_all["pipeline"] == pipeline) & (dm_all["dataset"] == dataset)] if len(dm_all) else pd.DataFrame()
            dm_wins = int(dm_sub["exp_lower_loss"].sum()) if len(dm_sub) and "exp_lower_loss" in dm_sub else 0
            dm_sig = int(dm_sub["significant_5pct"].sum()) if len(dm_sub) and "significant_5pct" in dm_sub else 0
            if len(exp):
                exp_row = exp.iloc[0]
                win_rows.append({
                    "pipeline": pipeline,
                    "dataset": dataset,
                    "asset": meta_map.get(dataset, {}).get("asset", ""),
                    "exp_f1": exp_row.get("f1"),
                    "best_f1_method": best_f1["method"],
                    "best_f1": best_f1["f1"],
                    "exp_auc": exp_row.get("auc"),
                    "best_auc_method": best_auc["method"],
                    "best_auc": best_auc["auc"],
                    "exp_best_by_f1": bool(exp_row.get("method") == best_f1["method"]),
                    "exp_best_by_auc": bool(exp_row.get("method") == best_auc["method"]),
                    "dm_lower_loss_wins_vs_3": dm_wins,
                    "dm_significant_5pct_wins_vs_3": dm_sig,
                    "outperform_summary": (
                        "exponential leads by F1 and DM" if (exp_row.get("method") == best_f1["method"] and dm_wins >= 2)
                        else "mixed; exponential not dominant"
                    ),
                })
    exp_wins = pd.DataFrame(win_rows)

    notes = pd.DataFrame([
        {"item": "scope", "detail": "US Treasury 1-month and 30-year yields from combine_TTM.xlsx sheet Ycurve."},
        {"item": "sample", "detail": f"Prepared slice {START_DATE} to {END_DATE}; train through {TRAIN_END_DATE}; OOS/test starts {TEST_START_DATE}; horizon {FORECAST_DAYS} windows."},
        {"item": "covid_label", "detail": f"External COVID crash label covers windows overlapping {CRASH_START_DATE} to {CRASH_END_DATE}."},
        {"item": "pipelines", "detail": "CNN autoencoder and standalone f-AnoGAN use raw yield GAF windows with external COVID labels; hybrid uses TadGAN anomaly scores -> GAF -> f-AnoGAN and backtest labels from TadGAN convention."},
        {"item": "dm_test", "detail": "DM table uses ranked anomaly score as probability-like forecast score and Brier loss vs crash label; positive DM statistic means exponential has lower loss than baseline."},
        {"item": "runtime", "detail": f"Reduced run: f-AnoGAN/TadGAN epochs={EPOCHS_SMALL}; CNN epochs={CNN_EPOCHS}; designed for GPU if torch CUDA is visible."},
    ])

    perf_all.to_csv(f"{SUMMARY_CSV_PREFIX}_performance_all.csv", index=False)
    perf_avg.to_csv(f"{SUMMARY_CSV_PREFIX}_performance_avg.csv", index=False)
    dm_all.to_csv(f"{SUMMARY_CSV_PREFIX}_dm_rank_score.csv", index=False)
    exp_wins.to_csv(f"{SUMMARY_CSV_PREFIX}_exponential_wins.csv", index=False)

    with pd.ExcelWriter(SUMMARY_XLSX, engine="openpyxl") as writer:
        meta.to_excel(writer, sheet_name="input_metadata", index=False)
        perf_all.to_excel(writer, sheet_name="performance_all", index=False)
        perf_avg.to_excel(writer, sheet_name="performance_avg", index=False)
        dm_all.to_excel(writer, sheet_name="dm_rank_score", index=False)
        exp_wins.to_excel(writer, sheet_name="exponential_wins", index=False)
        notes.to_excel(writer, sheet_name="methodology_notes", index=False)

    style_workbook(SUMMARY_XLSX)
    log(f"[SUMMARY] wrote {SUMMARY_XLSX}")
    print_tables(perf_all, perf_avg, dm_all, exp_wins)
    return perf_all, perf_avg, dm_all, exp_wins


def style_workbook(path: Path) -> None:
    if load_workbook is None or not path.exists():
        return
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for col_cells in ws.columns:
            max_len = 10
            letter = col_cells[0].column_letter
            for cell in col_cells[:200]:
                val = cell.value
                if val is not None:
                    max_len = max(max_len, min(60, len(str(val))))
            ws.column_dimensions[letter].width = max_len + 2
    wb.save(path)


def fmt_float(x: object) -> str:
    try:
        v = float(x)
        if not np.isfinite(v):
            return "nan"
        return f"{v:.4f}"
    except Exception:
        return str(x)


def print_tables(perf_all: pd.DataFrame, perf_avg: pd.DataFrame, dm_all: pd.DataFrame, exp_wins: pd.DataFrame) -> None:
    pd.set_option("display.width", 220)
    pd.set_option("display.max_columns", 30)
    if len(perf_all):
        cols = [c for c in ["pipeline", "dataset", "method", "accuracy", "precision", "recall", "f1", "auc", "n_test_total", "n_test_crash", "n_test_normal"] if c in perf_all.columns]
        log("\n===== PERFORMANCE (F1/AUC/precision) =====")
        log(perf_all[cols].sort_values(["pipeline", "dataset", "method"]).to_string(index=False, float_format=lambda x: f"{x:.4f}"))
    else:
        log("\n===== PERFORMANCE =====\nNo performance outputs found yet.")

    if len(perf_avg):
        cols = [c for c in ["pipeline", "method", "accuracy_mean", "precision_mean", "recall_mean", "f1_mean", "auc_mean", "n_rows"] if c in perf_avg.columns]
        log("\n===== AVERAGE ACROSS US1M + US30Y =====")
        log(perf_avg[cols].to_string(index=False, float_format=lambda x: f"{x:.4f}"))

    if len(dm_all):
        cols = [c for c in ["pipeline", "dataset", "comparison", "baseline_rank_brier_loss", "exponential_rank_brier_loss", "loss_diff_baseline_minus_exp", "dm_stat_positive_exp_better", "p_one_sided_exp_better", "exp_lower_loss", "significant_5pct"] if c in dm_all.columns]
        log("\n===== DM TEST (ranked anomaly score Brier loss; positive DM = exponential better) =====")
        log(dm_all[cols].sort_values(["pipeline", "dataset", "comparison"]).to_string(index=False, float_format=lambda x: f"{x:.4f}"))
    else:
        log("\n===== DM TEST =====\nNo DM outputs found yet.")

    if len(exp_wins):
        cols = [c for c in ["pipeline", "dataset", "exp_f1", "best_f1_method", "best_f1", "exp_auc", "best_auc_method", "best_auc", "dm_lower_loss_wins_vs_3", "dm_significant_5pct_wins_vs_3", "outperform_summary"] if c in exp_wins.columns]
        log("\n===== DOES EXPONENTIAL OUTPERFORM? =====")
        log(exp_wins[cols].sort_values(["pipeline", "dataset"]).to_string(index=False, float_format=lambda x: f"{x:.4f}"))


def main() -> None:
    parser = argparse.ArgumentParser(description="Run US bond 1M/30Y CNN, standalone f-AnoGAN, and hybrid comparisons.")
    parser.add_argument("--stage", choices=["prepare", "cnn", "fanogan", "hybrid", "summarize", "all"], default="all")
    parser.add_argument("--force", action="store_true", help="rerun stages even if outputs exist")
    args = parser.parse_args()

    t0 = time.time()
    log(f"[START] stage={args.stage} | base={BASE}")
    if args.stage in ("prepare", "all"):
        prepare_inputs()
    if args.stage in ("cnn", "all"):
        run_cnn(force=args.force)
    if args.stage in ("fanogan", "all"):
        run_fanogan(force=args.force)
    if args.stage in ("hybrid", "all"):
        run_hybrid(force=args.force)
    if args.stage in ("summarize", "all"):
        summarize()
    log(f"[DONE] elapsed_min={(time.time() - t0) / 60.0:.2f}")


if __name__ == "__main__":
    main()



